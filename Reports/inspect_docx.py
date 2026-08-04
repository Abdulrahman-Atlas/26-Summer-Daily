from docx import Document

doc = Document("Day1.docx")
print("Number of tables:", len(doc.tables))
for i, table in enumerate(doc.tables):
    print(f"Table {i}:")
    for r, row in enumerate(table.rows):
        print(f"  Row {r}:")
        for c, cell in enumerate(row.cells):
            print(f"    Cell {c}: {cell.text.strip()}")

